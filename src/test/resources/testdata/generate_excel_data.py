#!/usr/bin/env python3
"""
Generates login_data.xlsx for the mobile automation data-driven tests.
Run once during setup: python3 generate_excel_data.py
Requires: pip install openpyxl
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

OUTPUT = Path(__file__).parent / "login_data.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

LOGIN_DATA = [
    ("scenario", "username", "password", "expectSuccess"),
    ("Valid credentials - standard user", "testuser@example.com", "Test@1234", "true"),
    ("Invalid password",                  "testuser@example.com", "WrongPass",  "false"),
    ("Non-existent user",                 "nobody@nowhere.com",   "Test@1234",  "false"),
    ("Empty username",                    "",                     "Test@1234",  "false"),
    ("Empty password",                    "testuser@example.com", "",           "false"),
]


def write_sheet(wb, sheet_name, rows):
    ws = wb.create_sheet(sheet_name)
    for col_idx, header in enumerate(rows[0], start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(header) + 4, 20)

    for row_idx, row in enumerate(rows[1:], start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


wb = openpyxl.Workbook()
wb.remove(wb.active)  # Remove default sheet
write_sheet(wb, "LoginData", LOGIN_DATA)
wb.save(OUTPUT)
print(f"Created: {OUTPUT}")
